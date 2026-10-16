"""Monta o .docx (abre direto no Google Docs) a partir das páginas transcritas."""
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from modelo import Bloco, Pagina

ALINHAMENTO = {
    "esquerda": WD_ALIGN_PARAGRAPH.LEFT,
    "centro": WD_ALIGN_PARAGRAPH.CENTER,
    "direita": WD_ALIGN_PARAGRAPH.RIGHT,
}
MARCADORES_SIMPLES = {"", "•", "-", "*", "·", "–"}
RECUO = Cm(1.0)
LARGURA_MAX_IMAGEM = Cm(15)


def _escrever_trechos(paragrafo, trechos):
    for t in trechos:
        run = paragrafo.add_run(t.texto)
        run.bold = t.negrito or None
        run.italic = t.italico or None
        run.underline = t.sublinhado or None
        run.font.strike = t.riscado or None
        if t.destaque:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if t.cor and re.fullmatch(r"#?[0-9A-Fa-f]{6}", t.cor):
            run.font.color.rgb = RGBColor.from_string(t.cor.lstrip("#").upper())


def _nivel(bloco, maximo=3):
    return max(0, min(bloco.nivel, maximo))


def _item_lista(doc, bloco: Bloco):
    nivel = _nivel(bloco)
    if bloco.tipo == "lista_marcadores" and bloco.marcador.strip() in MARCADORES_SIMPLES:
        # Estilo de lista de verdade: o Google Docs importa como marcadores.
        estilo = "List Bullet" if nivel == 0 else f"List Bullet {min(nivel + 1, 3)}"
        p = doc.add_paragraph(style=estilo)
    else:
        # Numeração / marcadores próprios ("a)", "☐", "→") ficam como escritos,
        # assim cada lista começa do número que está no caderno.
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = RECUO * (nivel + 1)
        p.paragraph_format.first_line_indent = -Cm(0.6)
        p.add_run(f"{bloco.marcador.strip() or '•'}\t")
        p.paragraph_format.tab_stops.add_tab_stop(RECUO * (nivel + 1))
    _escrever_trechos(p, bloco.trechos)
    return p


def _tabela(doc, linhas):
    linhas = [l for l in linhas if l]
    if not linhas:
        return
    colunas = max(len(l) for l in linhas)
    tabela = doc.add_table(rows=len(linhas), cols=colunas)
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, linha in enumerate(linhas):
        for j, valor in enumerate(linha):
            celula = tabela.cell(i, j)
            celula.text = valor
            if i == 0:
                for run in celula.paragraphs[0].runs:
                    run.bold = True


def _recortar(imagem: Image.Image, caixa, folga=0.02):
    largura, altura = imagem.size
    x0, x1 = sorted((caixa.x0, caixa.x1))
    y0, y1 = sorted((caixa.y0, caixa.y1))
    x0, y0 = max(0.0, x0 - folga), max(0.0, y0 - folga)
    x1, y1 = min(1.0, x1 + folga), min(1.0, y1 + folga)
    if x1 - x0 < 0.03 or y1 - y0 < 0.03:
        return None
    recorte = imagem.crop((int(x0 * largura), int(y0 * altura), int(x1 * largura), int(y1 * altura)))
    recorte.thumbnail((1600, 1600))
    saida = io.BytesIO()
    recorte.convert("RGB").save(saida, "JPEG", quality=88)
    saida.seek(0)
    return saida


def _diagrama(doc, bloco: Bloco, imagem, modo):
    d = bloco.diagrama
    if d is None:
        return
    if modo in ("imagem", "ambos") and imagem is not None:
        recorte = _recortar(imagem, d.caixa)
        if recorte is not None:
            doc.add_picture(recorte, width=LARGURA_MAX_IMAGEM)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if modo in ("descricao", "ambos") or imagem is None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[Diagrama] {d.descricao}")
        r.italic = True
        r.font.size = Pt(9)
        if d.textos and modo == "descricao":
            r = p.add_run("\nTextos: " + " | ".join(d.textos))
            r.italic = True
            r.font.size = Pt(9)


def adicionar_pagina(doc, pagina: Pagina, imagem=None, modo_diagrama="ambos"):
    for bloco in pagina.blocos:
        if bloco.tipo == "titulo":
            p = doc.add_heading(level=max(1, min(bloco.nivel or 1, 3)))
            _escrever_trechos(p, bloco.trechos)
        elif bloco.tipo in ("lista_marcadores", "lista_numerada"):
            p = _item_lista(doc, bloco)
        elif bloco.tipo == "citacao":
            p = doc.add_paragraph(style="Quote")
            _escrever_trechos(p, bloco.trechos)
        elif bloco.tipo == "tabela":
            _tabela(doc, bloco.tabela or [])
            continue
        elif bloco.tipo == "diagrama":
            _diagrama(doc, bloco, imagem, modo_diagrama)
            continue
        elif bloco.tipo == "separador":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("⁂")
            continue
        else:
            p = doc.add_paragraph()
            if bloco.nivel:
                p.paragraph_format.left_indent = RECUO * _nivel(bloco)
            _escrever_trechos(p, bloco.trechos)
        p.alignment = ALINHAMENTO[bloco.alinhamento]


def novo_documento(titulo=None):
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)
    for secao in doc.sections:
        secao.left_margin = secao.right_margin = Cm(2.5)
    if titulo:
        doc.add_heading(titulo, level=0)
    return doc


def quebra_de_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def nota(doc, texto):
    """Nota pequena em cinza (origem da foto, trechos ilegíveis)."""
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
