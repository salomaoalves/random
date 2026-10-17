"""Builds the .docx (opens directly in Google Docs) from the transcribed pages."""
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from schema import Block, Page

ALIGNMENT = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}
PLAIN_MARKERS = {"", "•", "-", "*", "·", "–"}
INDENT = Cm(1.0)
MAX_IMAGE_WIDTH = Cm(15)


def _write_runs(paragraph, runs):
    for r in runs:
        run = paragraph.add_run(r.text)
        run.bold = r.bold or None
        run.italic = r.italic or None
        run.underline = r.underline or None
        run.font.strike = r.strike or None
        if r.highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if r.color and re.fullmatch(r"#?[0-9A-Fa-f]{6}", r.color):
            run.font.color.rgb = RGBColor.from_string(r.color.lstrip("#").upper())


def _level(block, maximum=3):
    return max(0, min(block.level, maximum))


def _list_item(doc, block: Block):
    level = _level(block)
    if block.type == "bullet_list" and block.marker.strip() in PLAIN_MARKERS:
        # Real list style: Google Docs imports it as bullets.
        style = "List Bullet" if level == 0 else f"List Bullet {min(level + 1, 3)}"
        p = doc.add_paragraph(style=style)
    else:
        # Numbering / custom markers ("a)", "☐", "→") are kept as written,
        # so each list starts from the number that is in the notebook.
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = INDENT * (level + 1)
        p.paragraph_format.first_line_indent = -Cm(0.6)
        p.add_run(f"{block.marker.strip() or '•'}\t")
        p.paragraph_format.tab_stops.add_tab_stop(INDENT * (level + 1))
    _write_runs(p, block.runs)
    return p


def _table(doc, rows):
    rows = [r for r in rows if r]
    if not rows:
        return
    columns = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def _crop(image: Image.Image, box, margin=0.02):
    width, height = image.size
    x0, x1 = sorted((box.x0, box.x1))
    y0, y1 = sorted((box.y0, box.y1))
    x0, y0 = max(0.0, x0 - margin), max(0.0, y0 - margin)
    x1, y1 = min(1.0, x1 + margin), min(1.0, y1 + margin)
    if x1 - x0 < 0.03 or y1 - y0 < 0.03:
        return None
    cropped = image.crop((int(x0 * width), int(y0 * height), int(x1 * width), int(y1 * height)))
    cropped.thumbnail((1600, 1600))
    buffer = io.BytesIO()
    cropped.convert("RGB").save(buffer, "JPEG", quality=88)
    buffer.seek(0)
    return buffer


def _diagram(doc, block: Block, image, mode):
    d = block.diagram
    if d is None:
        return
    if mode in ("image", "both") and image is not None:
        cropped = _crop(image, d.box)
        if cropped is not None:
            doc.add_picture(cropped, width=MAX_IMAGE_WIDTH)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if mode in ("description", "both") or image is None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[Diagram] {d.description}")
        r.italic = True
        r.font.size = Pt(9)
        if d.texts and mode == "description":
            r = p.add_run("\nTexts: " + " | ".join(d.texts))
            r.italic = True
            r.font.size = Pt(9)


def add_page(doc, page: Page, image=None, diagram_mode="both"):
    for block in page.blocks:
        if block.type == "heading":
            p = doc.add_heading(level=max(1, min(block.level or 1, 3)))
            _write_runs(p, block.runs)
        elif block.type in ("bullet_list", "numbered_list"):
            p = _list_item(doc, block)
        elif block.type == "quote":
            p = doc.add_paragraph(style="Quote")
            _write_runs(p, block.runs)
        elif block.type == "table":
            _table(doc, block.table or [])
            continue
        elif block.type == "diagram":
            _diagram(doc, block, image, diagram_mode)
            continue
        elif block.type == "separator":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("⁂")
            continue
        else:
            p = doc.add_paragraph()
            if block.level:
                p.paragraph_format.left_indent = INDENT * _level(block)
            _write_runs(p, block.runs)
        p.alignment = ALIGNMENT[block.alignment]


def new_document(title=None):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.5)
    if title:
        doc.add_heading(title, level=0)
    return doc


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def note(doc, text):
    """Small gray note (source photo, illegible parts)."""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
